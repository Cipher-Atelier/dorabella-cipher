from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
H=Path(__file__).resolve().parents[1]
s=json.loads((H/'results/summary.json').read_text());d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)

for style in d.styles:
 for b in list(style.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.06
d.styles['Title'].font.size=Pt(22);d.styles['Title'].paragraph_format.space_after=Pt(8)
d.styles['Heading 1'].font.size=Pt(13);d.styles['Heading 1'].paragraph_format.space_before=Pt(12);d.styles['Heading 1'].paragraph_format.space_after=Pt(5)
d.add_paragraph('Dorabella и сдвиг компонентов знака','Title')
d.add_paragraph('Результат нового вычислительного теста • 9 октября 2026 года')
p=d.add_paragraph();p.add_run('Расшифровка не получена. ').bold=True;p.add_run('Проверена новая для сохранённого исследования схема: направление знака и число его дуг относятся к разным позициям сообщения. Полный поиск завершён. Лучшие варианты не образуют связного текста и не выделяются убедительно среди контрольных последовательностей.')
d.add_heading('Что проверялось',1)
d.add_paragraph('У каждого из 87 знаков есть направление из восьми возможных и число дуг от одной до трёх. Эти два потока сдвигались друг относительно друга, затем снова соединялись в буквы. Проверены все глобальные циклические сдвиги и ограниченный общий сдвиг внутри каждой печатной строки, оба направления чтения и два порядка заполнения алфавитной таблицы.')
d.add_paragraph('На каждый вход приходится 133 632 набора параметров: 192 таблицы × 696 маршрутов. Повторяющиеся правила включены в этот счёт. Это не 133 632 независимые гипотезы. Каждый компонент каждого знака используется ровно один раз. Вставок, удалений, исправлений рукописи и подсказок предполагаемого текста нет.')
d.add_heading('Численный результат',1)
t=d.add_table(rows=1,cols=4);t.autofit=False
heads=['Транскрипция','Лучший балл S','Контроли не хуже','Описательный хвост']
for c,x in zip(t.rows[0].cells,heads):c.text=x
for name,label in [('hauer_real','Hauer'),('hart_real','Hartmeier')]:
 x=s['versions'][name];cells=t.add_row().cells
 for c,v in zip(cells,[label,f'{x["S"]:.6f}',f'{x["null_ge_S"]} из {x["n_null"]}',f'{x["plus_one_tail_S"]:.2f}']):c.text=v
for rowi,row in enumerate(t.rows):
 for col,c in enumerate(row.cells):
  c.width=Inches([1.35,1.5,1.6,2.2][col]);pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right','insideH','insideV']:
   q=OxmlElement('w:'+side);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');b.append(q)
  pr.append(b);m=OxmlElement('w:tcMar')
  for side in ['top','bottom','left','right']:
   q=OxmlElement('w:'+side);q.set(qn('w:w'),'90');q.set(qn('w:type'),'dxa');m.append(q)
  pr.append(m)
  if rowi==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');pr.append(sh)
   for r in c.paragraphs[0].runs:r.bold=True
  for p in c.paragraphs:p.paragraph_format.space_after=Pt(0)
d.add_paragraph().paragraph_format.space_after=Pt(0)
d.add_paragraph('Чем выше S, тем лучше оценка прежней языковой модели. Для каждой транскрипции заново созданы 199 контролей с теми же частотами знаков, направленными соседними парами и крайними знаками. На каждом выполнен тот же полный поиск.')
d.add_paragraph('Прибавка относительно лучшего правила без сдвига составила 0,614177 и 0,742512. В обеих версиях 30 из 199 контролей дали не меньшую прибавку; описательный хвост равен 0,155. Следовательно, само улучшение балла после добавления сдвига также не подтверждает механизм.')
d.add_page_break()
d.add_heading('Проверки вычислений',1)
for txt in [
'Все 14 известных сообщений после нового искусственного шифрования восстановлены полностью: 1218 из 1218 букв. Требование до запуска Dorabella: не менее 12 точных сообщений и 1194 верных букв. Это проверка чувствительности на прежних текстах, а не новая слепая валидация.',
'Основная серия содержит 414 входов: 14 известных сообщений, две транскрипции и 398 контролей. Всего проверено 55 323 648 наборов параметров. Три дополнительных полных поиска без отсечения совпали по максимумам, всем равным победителям и базовому варианту без сдвига. Повтор всей серии в новой папке также дал те же 414 результатов.',
'Проверены прямое и обратное преобразования, все таблицы и маршруты, сохранение свойств контролей и независимый подсчёт баллов. Независимо проверены все 414 результатов и 426 ключей с максимальным баллом. Максимальное численное расхождение — 5,33 × 10⁻¹⁵. Контрольные суммы использованных исходных файлов совпали.'
]:d.add_paragraph(txt)
d.add_heading('Границы вывода',1)
d.add_paragraph('Результат не поддерживает эту конкретную конечную схему. Он не исключает все дробящие шифры, произвольные таблицы, отдельные неизвестные сдвиги трёх строк или иные способы кодирования. Проверка обратного шифрования устанавливает математическую согласованность, но сама по себе не доказывает историческую правильность.')
d.add_paragraph('Обе транскрипции основаны на связанных чтениях одного шифра. Их результаты нельзя перемножать как независимые свидетельства. Число 0,16 — условный описательный ранг, а не вероятность того, что гипотеза верна. Свежие контроли не отменяют того, что новая гипотеза выбрана после предыдущих исследований. Ни одна буква Dorabella не объявляется установленной.')
d.add_heading('Что сохранено',1)
d.add_paragraph('Приложенный ZIP содержит зафиксированный протокол, исходные входы, прежнюю языковую модель, полный код, все 414 результатов, независимые проверки и запуск в новой папке для воспроизведения. Подробные правила, ограничения и команды находятся в README_REPRODUCE.txt. Полные бессвязные победители сохранены в результатах без искусственного превращения их в осмысленные фразы.')
d.add_paragraph('В этой проверке не использовались платные сканы, новые архивные документы или изображение British Library с ограничением на распространение. Следующий самостоятельный тест имеет смысл только при новом механизме с заранее определённой проверкой; повторное расширение этого сдвига ради красивого балла не обосновано.')
d.core_properties.title='Dorabella и сдвиг компонентов знака';d.core_properties.subject='Конечный воспроизводимый криптоаналитический тест';d.core_properties.author=''
d.save(H/'Dorabella_coordinate_lag_report_ru.docx')
