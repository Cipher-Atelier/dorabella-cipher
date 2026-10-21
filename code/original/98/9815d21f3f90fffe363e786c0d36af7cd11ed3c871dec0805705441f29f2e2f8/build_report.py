from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
H=Path(__file__).resolve().parents[1]
s=json.loads((H/'checks/results_verified.json').read_text());assert s['verified_cases']==94
checks=[json.loads((H/'checks'/('unpruned_'+c+'.json')).read_text()) for c in ['bailey_01','hauer_real','hart_real','hauer_real_pair_00']];assert all(x['status']=='passed' for x in checks)
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=D.styles[name];st.font.name='Liberation Sans';st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Normal'].paragraph_format.line_spacing=1.04
D.styles['Title'].font.size=Pt(23);D.styles['Heading 1'].font.size=Pt(14)
D.add_paragraph('Dorabella проверка независимого порядка дуг',style='Title')
D.add_paragraph('Раунд 14  •  5 октября 2026 года')
D.add_paragraph('Расшифровка не получена. Новая ограниченная модель полностью проверена на обеих транскрипциях и 78 контрольных последовательностях. Проверки на известных письмах прошли, но лучшие выводы для Dorabella остаются несвязными.')
D.add_heading('Что изменилось',level=1)
D.add_paragraph('Восемь троек ABC, DEF, GHI, KLM, NOP, QRS, TUW, XYZ расположены по кругу; разрешены восемь поворотов и два направления обхода. В каждой тройке разрешён любой из шести порядков букв по числу дуг, независимо от остальных направлений. I/J и U/V объединены. Все 87 знаков читаются в исходном порядке; исправлений, перестановок позиций и добавленных слов нет.')
D.add_paragraph('Это соседняя гипотеза, а не расширение всего раунда 13: внутренние перестановки стали свободнее, порядок самих троек ограничен. Полных ключей 16 × 6⁸ = 26 873 856 на вход; пересечение с прежней семьёй составляет 12 288 ключей. Круговая таблица здесь является проверяемым предположением, а не установленным ключом 1897 года.')
D.add_heading('Результаты',level=1)
t=D.add_table(rows=1,cols=4);t.autofit=False
width=[1.4,1.35,1.55,1.55]
for c,w in zip(t.columns,width):c.width=Inches(w)
for c,txt in zip(t.rows[0].cells,['Транскрипция','Лучший S87','Контроли не ниже','Хвост с +1']):c.text=txt
for name,label in [('hauer_real','Hauer'),('hart_real','Hartmeier')]:
 r=s[name];cs=t.add_row().cells
 for c,txt in zip(cs,[label,f"{r['real']['S']:.6f}",f"{r['controls_no_lower']} из 39",f"{r['descriptive_tail']:.3f}"]):c.text=txt
for ri,row in enumerate(t.rows):
 for cell in row.cells:
  tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');borders.append(z)
  tcPr.append(borders)
  if ri==0:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');tcPr.append(shade)
   for p in cell.paragraphs:
    for r in p.runs:r.bold=True
p_tail=D.add_paragraph('Хвосты рассчитаны отдельно: (1 + число контролей с не меньшей оценкой) / 40. Это описательное сравнение с условной нулевой моделью, не вероятность правильности ключа. Контроли сохраняют частоты знаков, направленные соседние пары и крайние знаки. Они уже использовались; выбор новой семьи сделан после предыдущих опытов.')
p_tail.paragraph_format.space_before=Pt(6)
D.add_paragraph('Все 14 прежних контрольных писем, заново зашифрованных новой моделью, восстановлены полностью: 1218 из 1218 букв. Настоящий ключ каждого примера входит в множество точных максимумов. Это проверка чувствительности на регрессионных текстах, а не новая независимая валидация.')
D.add_paragraph('Обе транскрипции выбрали ключ с каноническим ID 19301530. У него ровно один равнозначный полный ключ, ID 19305418: различие относится к отсутствующим знакам и не меняет полученный текст. Найденные отдельные английские фрагменты не складываются в сообщение. Никакие буквы Dorabella не установлены.')
D.add_page_break()
D.add_heading('Проверки и границы вывода',level=1)
D.add_paragraph(f"Завершены все 94 входа: {s['char_evaluations']:,} символьных оценок и {s['word_evaluations']:,} полных словарных оценок. Остальные словарные расчёты отсечены безопасной верхней границей. На каждом ключе учитывались все 87 знаков; отложенного участка нет.".replace(',',' '))
D.add_paragraph('Четыре полных повторения без отсечения совпали по максимумам и всем точным ничьим: ещё 107 495 424 полных оценки C и W. Независимо проверены итоговые ключи, все ничьи, обратное шифрование, цепочки контрольных точек и неизменность зафиксированных файлов. До запуска прошли 132 независимые проверки и 12 сравнений с прямым перебором по 1000 ключей.')
D.add_paragraph('Вывод относится только к этой конечной семье и неизменённому критерию S87 = C87 + 0.5 W87. Он не исключает другие подстановки, иной порядок троек, шифры с памятью, перестановки позиций, другую орфографию или ошибку языкового критерия. Новых исторических источников и независимого подтверждения чтения этот опыт не добавляет.')
D.add_heading('Как повторить',level=1)
D.add_paragraph('В архиве находятся входы, модель, зафиксированный протокол, код, все 1504 контрольные точки, 94 результата и проверки. Установите версии из requirements.txt. Из каталога пакета выполните python code/reproduce.py --output /путь/к/новому/каталогу. Команда создаёт новую рабочую папку, заново проходит контрольные письма и сравнивает результаты с эталонами. Для одного реального случая добавьте --case hauer_real; 14 проверок чувствительности всё равно выполняются заново. Исходный run.py не следует использовать поверх непроверенных старых результатов.')
D.add_paragraph('Протокол зафиксирован 5 октября 2026 года в 13:44:24 UTC, до реальных вычислений. SHA256: ae5d77723a64e36ecef03edf596339b6da027173af0bc601df7fece0bb756996.')
D.add_heading('Источники данных',level=1)
D.add_paragraph('Hauer et al. Dorabella Cipher as Musical Inspiration, 2021, Figure 2: https://aclanthology.org/2021.smp-1.5.pdf. Daniel Hartmeier, Clues (or red herrings) to the Dorabella Cipher: https://www.benzedrine.ch/dorabella.html. Сохранённый воспроизводимый пакет раунда 13 служит источником модели, регрессионных текстов и контрольных последовательностей; его хеш указан в protocol.json.')
for p in D.paragraphs:
 if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
for root in (D._element,D.styles.element):
 for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
D.core_properties.title='Dorabella проверка независимого порядка дуг';D.core_properties.author=''
D.save(H/'Dorabella_round14_report_ru.docx')
